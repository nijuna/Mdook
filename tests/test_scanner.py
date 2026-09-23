from pathlib import Path

import docx
import pymupdf
import pytest
from ebooklib import epub

from mdook.core.scanner import (
    _clean_title,
    format_file_size,
    scan_directory,
    sniff_book_metadata,
)


def test_format_file_size() -> None:
    assert format_file_size(500) == "500 B"
    assert format_file_size(1024) == "1.0 KB"
    assert format_file_size(1536) == "1.5 KB"
    assert format_file_size(1024 * 1024 * 5) == "5.0 MB"
    assert format_file_size(1024 * 1024 * 1024 * 2) == "2.00 GB"


def test_clean_title() -> None:
    expected = "The Hound of the Baskervilles"
    assert _clean_title(expected, fallback="default") == expected
    assert _clean_title(None, fallback="default") == "default"
    assert _clean_title("", fallback="default") == "default"
    assert _clean_title("   ", fallback="default") == "default"
    assert _clean_title("C:\\Users\\Book\\title.htm", fallback="default") == "default"
    assert _clean_title("file%20path%20encoded", fallback="default") == "default"
    assert _clean_title("A" * 151, fallback="default") == "default"
    assert (
        _clean_title("silliman_chinese.qxd", fallback="Electronic Revolution")
        == "Electronic Revolution"
    )
    assert _clean_title("book_layout.indd", fallback="Real Book") == "Real Book"
    assert (
        _clean_title("The Body Keeps the Score - PDFDrive.com", fallback="default")
        == "The Body Keeps the Score"
    )
    assert _clean_title("Nietzsche - Z-Library", fallback="default") == "Nietzsche"


def test_scan_directory_errors(tmp_path: Path) -> None:
    non_existent = tmp_path / "does_not_exist"
    with pytest.raises(FileNotFoundError):
        scan_directory(non_existent)

    regular_file = tmp_path / "file.txt"
    regular_file.write_text("hello")
    with pytest.raises(NotADirectoryError):
        scan_directory(regular_file)


def test_scan_directory_empty(tmp_path: Path) -> None:
    results = scan_directory(tmp_path)
    assert results == []


def test_scan_directory_filters_unsupported_and_hidden(tmp_path: Path) -> None:
    (tmp_path / "notes.txt").write_text("text")
    (tmp_path / "photo.png").write_bytes(b"image")
    (tmp_path / ".hidden.pdf").write_bytes(b"%PDF-1.4 dummy")
    (tmp_path / "~$temp.docx").write_bytes(b"dummy")

    results = scan_directory(tmp_path)
    assert results == []


def test_scan_directory_with_real_formats(tmp_path: Path) -> None:
    # 1. Create a sample PDF
    pdf_path = tmp_path / "sherlock.pdf"
    doc = pymupdf.open()
    doc.set_metadata({"title": "A Study in Scarlet", "author": "Arthur Conan Doyle"})
    page = doc.new_page(width=400, height=600)
    page.insert_text((50, 100), "Chapter 1: Mr. Sherlock Holmes")
    doc.save(str(pdf_path))
    doc.close()

    # 2. Create a sample EPUB
    epub_path = tmp_path / "dracula.epub"
    book = epub.EpubBook()
    book.set_identifier("test-isbn-1234")
    book.set_title("Dracula")
    book.add_author("Bram Stoker")
    ch = epub.EpubHtml(title="Chapter 1", file_name="ch01.xhtml", lang="en")
    ch.content = "<html><body><h1>Jonathan Harker's Journal</h1></body></html>"
    book.add_item(ch)
    book.spine = [ch]
    epub.write_epub(str(epub_path), book)

    # 3. Create a sample DOCX
    docx_path = tmp_path / "frankenstein.docx"
    doc_docx = docx.Document()
    doc_docx.core_properties.title = "Frankenstein"
    doc_docx.core_properties.author = "Mary Shelley"
    doc_docx.add_heading("Letter 1", level=1)
    doc_docx.add_paragraph("To Mrs. Saville, England.")
    doc_docx.save(str(docx_path))

    results = scan_directory(tmp_path)
    assert len(results) == 3

    # Results should be sorted alphabetically by filename:
    # dracula.epub, frankenstein.docx, sherlock.pdf
    dracula = results[0]
    assert dracula.filename == "dracula.epub"
    assert dracula.format == "epub"
    assert dracula.title == "Dracula"
    assert dracula.author == "Bram Stoker"
    assert dracula.size_bytes > 0
    assert "KB" in dracula.formatted_size or "B" in dracula.formatted_size

    frankenstein = results[1]
    assert frankenstein.filename == "frankenstein.docx"
    assert frankenstein.format == "docx"
    assert frankenstein.title == "Frankenstein"
    assert frankenstein.author == "Mary Shelley"

    sherlock = results[2]
    assert sherlock.filename == "sherlock.pdf"
    assert sherlock.format == "pdf"
    assert sherlock.title == "A Study in Scarlet"
    assert sherlock.author == "Arthur Conan Doyle"


def test_scan_directory_recursive(tmp_path: Path) -> None:
    sub_dir = tmp_path / "classics" / "gothic"
    sub_dir.mkdir(parents=True)

    pdf_root = tmp_path / "root_book.pdf"
    doc = pymupdf.open()
    doc.set_metadata({"title": "Root Book", "author": "Root Author"})
    doc.new_page()
    doc.save(str(pdf_root))
    doc.close()

    pdf_nested = sub_dir / "nested_book.pdf"
    doc2 = pymupdf.open()
    doc2.set_metadata({"title": "Nested Book", "author": "Nested Author"})
    doc2.new_page()
    doc2.save(str(pdf_nested))
    doc2.close()

    # In ignored dir (.git)
    git_dir = tmp_path / ".git"
    git_dir.mkdir()
    pdf_ignored = git_dir / "ignored.pdf"
    doc3 = pymupdf.open()
    doc3.new_page()
    doc3.save(str(pdf_ignored))
    doc3.close()

    # Non-recursive: only finds root
    flat_results = scan_directory(tmp_path, recursive=False)
    assert len(flat_results) == 1
    assert flat_results[0].filename == "root_book.pdf"

    # Recursive: finds root and nested, ignores .git
    rec_results = scan_directory(tmp_path, recursive=True)
    assert len(rec_results) == 2
    filenames = {b.filename for b in rec_results}
    assert filenames == {"root_book.pdf", "nested_book.pdf"}


def test_sniff_book_metadata_corrupt_file(tmp_path: Path) -> None:
    corrupt_pdf = tmp_path / "corrupt.pdf"
    corrupt_pdf.write_bytes(b"corrupted binary data")

    title, author = sniff_book_metadata(corrupt_pdf)
    assert title == "corrupt"
    assert author == "Unknown"

import io
import zipfile
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from mdook.core.formats.docx import parse_docx
from mdook.core.models import (
    BlockQuote,
    CodeBlock,
    ImageRef,
    ListData,
    Paragraph,
    TableData,
)
from mdook.core.pipeline import convert

SAMPLE_PNG_BYTES = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00"
    b"\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _create_sample_docx(
    path: Path,
    title: str = "Frankenstein",
    author: str = "Mary Shelley",
    is_technical: bool = False,
) -> Path:
    doc = docx.Document()
    doc.core_properties.title = title
    doc.core_properties.author = author
    doc.core_properties.comments = "A classic manuscript."

    # Front Matter: Preface
    doc.add_heading("Preface", level=1)
    doc.add_paragraph("The event on which this fiction is founded has been supposed possible.")

    # Chapter 1
    if is_technical:
        doc.add_heading("Chapter 1: Technical Architecture", level=1)
        doc.add_heading("1.1 System Overview", level=2)
        doc.add_paragraph("This technical specification outlines the infrastructure design.")

        # Code Block
        code_p = doc.add_paragraph("import os\ndef start_worker():\n    return os.getpid()")
        for r in code_p.runs:
            r.font.name = "Courier New"

        # Table
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Service"
        tbl.cell(0, 1).text = "Latency"
        tbl.cell(1, 0).text = "Gateway"
        tbl.cell(1, 1).text = "5ms"

        doc.add_heading("1.2 Benchmark Results", level=2)
        doc.add_paragraph("Throughput exceeded 10,000 requests per second.")
    else:
        doc.add_heading("Chapter 1: The Experiment", level=1)
        p_body = doc.add_paragraph("I am by birth a ")
        r_bold = p_body.add_run("Genevese")
        r_bold.bold = True
        p_body.add_run(", and my family is one of the most ")
        r_ital = p_body.add_run("distinguished")
        r_ital.italic = True
        p_body.add_run(" of that republic.")

        # Quote
        doc.add_paragraph(
            "To examine the causes of life, we must first have recourse to death.", style="Quote"
        )

        # Lists
        doc.add_paragraph("First observation", style="List Bullet")
        doc.add_paragraph("Second observation", style="List Bullet")
        doc.add_paragraph("Primary step", style="List Number")
        doc.add_paragraph("Secondary step", style="List Number")

        # Image
        img_stream = io.BytesIO(SAMPLE_PNG_BYTES)
        doc.add_picture(img_stream)

        # Footnote reference in paragraph
        p_fn = doc.add_paragraph("Natural philosophy is the genius that has regulated my fate")
        r_fn = p_fn.add_run()
        fn_ref = OxmlElement("w:footnoteReference")
        fn_ref.set(qn("w:id"), "1")
        r_fn._r.append(fn_ref)

    # Back Matter: Appendix
    doc.add_heading("Appendix", level=1)
    doc.add_paragraph("Historical and geographical context of Geneva.")

    doc.save(str(path))

    # Inject word/footnotes.xml if literature
    if not is_technical:
        fn_xml = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">\n'
            b'  <w:footnote w:id="1">\n'
            b"    <w:p><w:r><w:t>Referring to early natural sciences.</w:t></w:r></w:p>\n"
            b"  </w:footnote>\n"
            b"</w:footnotes>"
        )
        with zipfile.ZipFile(path, "a") as zf:
            zf.writestr("word/footnotes.xml", fn_xml)

    return path


def test_docx_metadata_and_structure(tmp_path: Path) -> None:
    docx_file = tmp_path / "frankenstein.docx"
    _create_sample_docx(docx_file, is_technical=False)

    manifest, tree = parse_docx(docx_file, profile_override="auto")

    assert manifest.title == "Frankenstein"
    assert manifest.author == "Mary Shelley"
    assert manifest.profile == "literature"
    assert manifest.needs_ocr is False
    assert manifest.total_pages >= 1

    # Check front matter
    assert len(tree.front_matter) >= 1
    assert any(sec.title == "Preface" for sec in tree.front_matter)

    # Check chapters
    assert len(tree.chapters) == 1
    ch = tree.chapters[0]
    assert "Chapter 1" in ch.title

    # Check back matter
    assert len(tree.back_matter) >= 1
    assert any(sec.title == "Appendix" for sec in tree.back_matter)


def test_docx_content_types(tmp_path: Path) -> None:
    docx_file = tmp_path / "content_test.docx"
    _create_sample_docx(docx_file, is_technical=False)

    manifest, tree = parse_docx(docx_file)
    ch = tree.chapters[0]
    all_content = [item for sec in ch.sections for item in sec.content]

    # Paragraph with bold and italic
    paragraphs = [c for c in all_content if isinstance(c, Paragraph)]
    assert any("**Genevese**" in p.text for p in paragraphs)
    assert any("*distinguished*" in p.text for p in paragraphs)

    # BlockQuote
    quotes = [c for c in all_content if isinstance(c, BlockQuote)]
    assert len(quotes) >= 1
    assert any("causes of life" in line for line in quotes[0].lines)

    # ListData
    lists = [c for c in all_content if isinstance(c, ListData)]
    assert len(lists) >= 1

    # Image
    images = [c for c in all_content if isinstance(c, ImageRef)]
    assert len(images) >= 1
    assert Path(images[0].source_path).exists()

    # Footnote
    assert len(ch.footnotes) == 1
    assert ch.footnotes[0].marker == "1"
    assert "early natural sciences" in ch.footnotes[0].text


def test_docx_technical_profile_auto_detection(tmp_path: Path) -> None:
    docx_file = tmp_path / "tech_spec.docx"
    _create_sample_docx(docx_file, is_technical=True)

    manifest, tree = parse_docx(docx_file, profile_override="auto")
    assert manifest.profile == "technical"

    ch = tree.chapters[0]
    all_content = [item for sec in ch.sections for item in sec.content]

    # CodeBlock
    codes = [c for c in all_content if isinstance(c, CodeBlock)]
    assert len(codes) >= 1
    assert any("start_worker" in line for line in codes[0].lines)

    # TableData
    tables = [c for c in all_content if isinstance(c, TableData)]
    assert len(tables) >= 1
    assert any(row[0] == "Service" for row in tables[0].cells)


def test_docx_end_to_end_conversion(tmp_path: Path) -> None:
    docx_file = tmp_path / "full_novel.docx"
    _create_sample_docx(docx_file, is_technical=False)

    vault_dir = tmp_path / "vaults"
    result = convert(docx_file, vault_dir)

    assert result.success is True
    assert result.output_dir.exists()
    assert result.chapters == 1
    assert result.footnotes == 1

    md_files = list(result.output_dir.glob("*.md"))
    filenames = {f.name for f in md_files}
    assert any("Index.md" in fn for fn in filenames)
    assert any("01 - " in fn for fn in filenames)
    assert "00 - Front Matter.md" in filenames
    assert "Appendix.md" in filenames

    # Check attachments
    attachments_dir = result.output_dir / "attachments"
    assert attachments_dir.exists()
    assert len(list(attachments_dir.glob("*.png"))) >= 1

    # Check footnote rendering
    chapter_file = next(f for f in md_files if "01 - " in f.name)
    content = chapter_file.read_text()
    assert "[^1]" in content
    assert "[^1]:" in content
